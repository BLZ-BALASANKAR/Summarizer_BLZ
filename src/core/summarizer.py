"""
Unified Summarizer Engine
Consolidates extractive and abstractive summarization logic.
"""

from typing import Dict, Any
import time
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from extractive_summarizers import ExtractiveSummarizers
from abstractive_summarizers import AbstractiveSummarizers
from evaluation_metrics import SummarizationEvaluator
from src.utils.logger import setup_logger

logger = setup_logger("SummarizerEngine")


class SummarizerEngine:
    def __init__(self, config: dict):
        self.config = config
        self.use_gpu = config.get("processing", {}).get("use_gpu", False)

        logger.info("Initializing Summarizers...")
        self.extractive = ExtractiveSummarizers()
        self.abstractive = AbstractiveSummarizers(use_gpu=self.use_gpu)
        self.evaluator = SummarizationEvaluator()

    def run(self, text: str) -> Dict[str, Any]:
        """Run all configured summarizers on the text."""
        start_time = time.time()
        results = {
            "input_text": text,
            "extractive_summaries": {},
            "abstractive_summaries": {},
            "evaluation": {},
            "metadata": {},
        }

        # 1. Extractive
        algo_list = self.config.get("extractive", {}).get("algorithms", [])
        if algo_list:
            logger.info(f"Running Extractive: {algo_list}")
            # Note: The existing ExtractiveSummarizers.summarize_all runs ALL by default
            # We might want to filter, but for now let's keep it simple
            results["extractive_summaries"] = self.extractive.summarize_all(text)

        # 2. Abstractive
        model_list = self.config.get("abstractive", {}).get("default_to_run", [])
        params = self.config.get("abstractive", {}).get("generation", {})

        if model_list:
            logger.info(f"Running Abstractive: {model_list}")
            for model in model_list:
                try:
                    summary = self._run_abstractive_model(model, text, params)
                    results["abstractive_summaries"][model] = summary
                except Exception as e:
                    logger.error(f"Model {model} failed: {e}")
                    results["abstractive_summaries"][model] = {"error": str(e)}

        # 3. Evaluation
        logger.info("Evaluating Summaries...")
        results["evaluation"] = self._evaluate_all(text, results)

        # Metadata
        results["metadata"] = {
            "execution_time": time.time() - start_time,
            "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
        }

        return results

    def _run_abstractive_model(self, model: str, text: str, params: dict):
        start = time.time()
        max_l = params.get("max_length", 150)
        min_l = params.get("min_length", 50)

        # Dispatch based on model name (simplified mapping)
        summary_text = ""
        if "bart" in model:
            res = self.abstractive.bart_summarize(text, max_l, min_l)
            summary_text = res["summary"]
        elif "t5" in model:
            if "flan" in model:
                variant = "large" if "large" in model else "base"
                res = self.abstractive.flan_t5_summarize(text, max_l, min_l, variant)
            else:
                res = self.abstractive.t5_summarize(text, max_l, min_l)
            summary_text = res["summary"]
        elif "pegasus" in model:
            res = self.abstractive.pegasus_summarize(text, max_l, min_l)
            summary_text = res["summary"]
        else:
            # Fallback or error
            raise ValueError(f"Unknown model architecture: {model}")

        return {"summary": summary_text, "execution_time": time.time() - start}

    def _evaluate_all(self, original_text: str, results: dict) -> dict:
        evals = {"extractive": {}, "abstractive": {}}

        # Extractive
        for algo, data in results.get("extractive_summaries", {}).items():
            if "summary" in data:
                # Use original text as reference for ROUGE/BLEU/METEOR calculation
                evals["extractive"][algo] = self.evaluator.evaluate_summary(
                    original_text,
                    data["summary"],
                    reference_summary=original_text,  # Enable ROUGE/BLEU/METEOR
                )
                # Add source mapping for Traceability
                evals["extractive"][algo]["source_mapping"] = (
                    self.evaluator.get_source_mapping(original_text, data["summary"])
                )

        # Abstractive
        for model, data in results.get("abstractive_summaries", {}).items():
            if "summary" in data:
                # Use original text as reference for ROUGE/BLEU/METEOR calculation
                evals["abstractive"][model] = self.evaluator.evaluate_summary(
                    original_text,
                    data["summary"],
                    reference_summary=original_text,  # Enable ROUGE/BLEU/METEOR
                )
                # Add source mapping for abstractive models too (shows which source influenced the AI)
                evals["abstractive"][model]["source_mapping"] = (
                    self.evaluator.get_source_mapping(original_text, data["summary"])
                )

        return evals

    def save_report(
        self, results: dict, filename: str = "full_summarization_report.xlsx"
    ):
        """Generate Excel Report (moved from excel_demo.py)"""
        logger.info(f"Generating Excel report: {filename}")

        # Helper to flat metrics
        def flatten_metrics(metrics):
            if not metrics:
                return {}
            flat = {
                "Compression %": metrics.get("compression_ratio", 0),
                "Density %": metrics.get("density", 0),
                "Cosine Sim": metrics.get("text_similarity", 0),
            }
            # ROUGE scores
            if "rouge" in metrics and "error" not in metrics["rouge"]:
                flat["ROUGE-1 F1"] = metrics["rouge"].get("rouge1_f", 0)
                flat["ROUGE-2 F1"] = metrics["rouge"].get("rouge2_f", 0)
                flat["ROUGE-L F1"] = metrics["rouge"].get("rougeL_f", 0)
            # BLEU scores
            if "bleu" in metrics and "error" not in metrics["bleu"]:
                flat["BLEU Score"] = metrics["bleu"].get("bleu_overall", 0)
            # METEOR score
            if "meteor" in metrics and "error" not in metrics["meteor"]:
                flat["METEOR"] = metrics["meteor"].get("meteor", 0)
            return flat

        all_rows = []

        # Process Extractive
        evals = results.get("evaluation", {}).get("extractive", {})
        for algo, data in results.get("extractive_summaries", {}).items():
            if "error" in data:
                continue
            row = {
                "Method": algo,
                "Type": "Extractive",
                "Summary": data["summary"],
                "Time (s)": data.get("execution_time", 0),
            }
            row.update(flatten_metrics(evals.get(algo, {})))
            all_rows.append(row)

        # Process Abstractive
        evals = results.get("evaluation", {}).get("abstractive", {})
        for model, data in results.get("abstractive_summaries", {}).items():
            if "error" in data:
                continue
            row = {
                "Method": model,
                "Type": "Abstractive",
                "Summary": data["summary"],
                "Time (s)": data.get("execution_time", 0),
            }
            row.update(flatten_metrics(evals.get(model, {})))
            all_rows.append(row)

        # Create DataFrame
        if all_rows:
            df = pd.DataFrame(all_rows)
            # Reorder columns preference
            cols = [
                "Method",
                "Type",
                "Summary",
                "Time (s)",
                "Compression %",
                "Density %",
                "Cosine Sim",
                "ROUGE-1 F1",
                "ROUGE-2 F1",
                "ROUGE-L F1",
                "BLEU Score",
                "METEOR",
            ]
            # Filter cols that exist
            final_cols = [c for c in cols if c in df.columns] + [
                c for c in df.columns if c not in cols
            ]
            df = df[final_cols]

            with pd.ExcelWriter(filename, engine="openpyxl") as writer:
                df.to_excel(writer, sheet_name="All Results", index=False)

            # Formatting (Optional but nice)
            self._format_excel(filename)

    def _format_excel(self, filename):
        try:
            wb = load_workbook(filename)
            ws = wb["All Results"]
            # Simple header styling
            header_font = Font(bold=True, color="FFFFFF")
            fill = PatternFill(
                start_color="4472C4", end_color="4472C4", fill_type="solid"
            )
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = fill
            wb.save(filename)
        except Exception as e:
            logger.warning(f"Excel formatting failed: {e}")
